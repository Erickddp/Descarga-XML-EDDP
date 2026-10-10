# -*- coding: utf-8 -*-
"""Exportación (Excel, CSV, portapapeles) e índice de lo ya leído.

El índice guarda en un CSV lo que ya se procesó, con la fecha de modificación
del archivo. Así la segunda vez no se vuelven a abrir miles de XML, y nunca hay
que volver a pedirle nada al SAT.
"""
import csv
import json
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from .lector import COLUMNAS, NUMERICAS, Renglon


# ---------------------------------------------------------------- portapapeles
def a_tsv(filas: List[Renglon], con_encabezado: bool = True) -> str:
    """Texto separado por tabuladores: se pega directo en Excel con Ctrl+V."""
    def limpiar(v) -> str:
        s = "" if v is None else str(v)
        return s.replace("\t", " ").replace("\r", " ").replace("\n", " ")

    lineas = []
    if con_encabezado:
        lineas.append("\t".join(COLUMNAS))
    for r in filas:
        lineas.append("\t".join(limpiar(v) for v in r.como_lista()))
    return "\n".join(lineas)


# ---------------------------------------------------------------- CSV
def a_csv(filas: List[Renglon], destino: str) -> str:
    d = Path(destino)
    d.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig para que Excel en Windows respete los acentos
    with open(d, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNAS)
        for r in filas:
            w.writerow(r.como_lista())
    return str(d)


# ---------------------------------------------------------------- Excel
def a_excel(filas: List[Renglon], destino: str, hoja: str = "CFDI") -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    d = Path(destino)
    d.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = hoja

    ws.append(COLUMNAS)
    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = encabezado
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r in filas:
        ws.append(r.como_lista())

    # formato de número y ancho de columna
    for i, col in enumerate(COLUMNAS, start=1):
        letra = get_column_letter(i)
        if col in NUMERICAS:
            fmt = "#,##0" if col == "Conceptos" else "#,##0.00"
            for celda in ws[letra][1:]:
                celda.number_format = fmt
        largos = [len(col)] + [len(str(r.datos.get(col, ""))) for r in filas[:400]]
        ws.column_dimensions[letra].width = min(max(largos) + 2, 42)

    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(filas) + 1}"

    _hoja_resumen(wb, filas)
    wb.save(d)
    return str(d)


def _hoja_resumen(wb, filas: List[Renglon]) -> None:
    from openpyxl.styles import Font
    ws = wb.create_sheet("Resumen")
    ws.append(["Mes", "Tipo", "Comprobantes", "SubTotal", "IVA Trasladado",
               "IVA Retenido", "ISR Retenido", "Total"])
    for c in ws[1]:
        c.font = Font(bold=True)

    agrupado: Dict[Tuple[str, str], List[float]] = {}
    for r in filas:
        k = (f"{r.datos.get('Año','')}-{r.datos.get('Mes','')}", str(r.datos.get("Tipo", "")))
        a = agrupado.setdefault(k, [0, 0.0, 0.0, 0.0, 0.0, 0.0])
        a[0] += 1
        for i, col in enumerate(["SubTotal", "IVA Trasladado", "IVA Retenido",
                                 "ISR Retenido", "Total"], start=1):
            a[i] += float(r.datos.get(col) or 0)

    for (mes, tipo), a in sorted(agrupado.items()):
        ws.append([mes, tipo, a[0], round(a[1], 2), round(a[2], 2),
                   round(a[3], 2), round(a[4], 2), round(a[5], 2)])
    for fila in ws.iter_rows(min_row=2, min_col=4):
        for celda in fila:
            celda.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12


# ---------------------------------------------------------------- índice
class Indice:
    """Recuerda qué XML ya se leyeron (ruta + fecha de modificación + tamaño).

    No guarda el contenido: sirve para saltarse el re-parseo y para saber qué
    hay descargado sin volver a preguntarle al SAT.
    """

    CAMPOS = ["archivo", "mtime", "tamano", "uuid", "tipo", "fecha", "total"]

    def __init__(self, ruta: str):
        self.ruta = Path(ruta)
        self.datos: Dict[str, Dict] = {}
        self.cargar()

    def cargar(self) -> None:
        self.datos = {}
        if not self.ruta.exists():
            return
        try:
            with open(self.ruta, newline="", encoding="utf-8") as f:
                for fila in csv.DictReader(f):
                    if fila.get("archivo"):
                        self.datos[fila["archivo"]] = fila
        except (OSError, csv.Error):
            self.datos = {}

    def guardar(self) -> None:
        self.ruta.parent.mkdir(parents=True, exist_ok=True)
        with open(self.ruta, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=self.CAMPOS)
            w.writeheader()
            for v in self.datos.values():
                w.writerow({k: v.get(k, "") for k in self.CAMPOS})

    @staticmethod
    def _firma(p: Path) -> Tuple[str, str]:
        st = p.stat()
        return str(int(st.st_mtime)), str(st.st_size)

    def esta_al_dia(self, ruta) -> bool:
        p = Path(ruta)
        v = self.datos.get(str(p))
        if not v or not p.exists():
            return False
        mtime, tam = self._firma(p)
        return v.get("mtime") == mtime and v.get("tamano") == tam

    def registrar(self, renglon: Renglon) -> None:
        p = Path(str(renglon.datos.get("Archivo", "")))
        if not p.exists():
            return
        mtime, tam = self._firma(p)
        self.datos[str(p)] = {
            "archivo": str(p), "mtime": mtime, "tamano": tam,
            "uuid": renglon.datos.get("UUID", ""),
            "tipo": renglon.datos.get("Tipo", ""),
            "fecha": renglon.datos.get("Fecha", ""),
            "total": renglon.datos.get("Total", ""),
        }

    def olvidar_inexistentes(self) -> int:
        fuera = [k for k in self.datos if not Path(k).exists()]
        for k in fuera:
            del self.datos[k]
        return len(fuera)

    def __len__(self) -> int:
        return len(self.datos)


# ---------------------------------------------------------------- caché de renglones
VERSION_CACHE = 1        # súbela si cambia lo que lector.leer_xml_filas saca de un XML


class CacheFilas:
    """Renglones ya interpretados de una carpeta, para no volver a abrir cada XML
    (en Google Drive o con antivirus, abrir miles de archivos es lo más lento).
    Un archivo por carpeta leída, junto al índice: cache_filas/<hash>.json.
    Cada entrada vale mientras el XML conserve fecha de modificación y tamaño."""

    def __init__(self, ruta_indice, carpeta: str):
        import hashlib
        clave = hashlib.sha1(str(Path(carpeta).resolve()).lower().encode("utf-8")).hexdigest()[:16]
        self.ruta = Path(ruta_indice).parent / "cache_filas" / f"{clave}.json"
        self.datos: Dict[str, list] = {}
        try:
            leido = json.loads(self.ruta.read_text(encoding="utf-8"))
            if leido.get("version") == VERSION_CACHE:
                self.datos = leido.get("archivos") or {}
        except (OSError, ValueError, AttributeError):
            self.datos = {}

    def filas_de(self, ruta: str, firma) -> Optional[List[Dict]]:
        v = self.datos.get(ruta)
        if v and v[0] == firma[0] and v[1] == firma[1]:
            return v[2]
        return None

    def guardar(self, vigentes: Dict[str, list]) -> None:
        """Deja solo lo que sigue existiendo en la carpeta."""
        self.datos = vigentes
        try:
            self.ruta.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.ruta.with_suffix(".tmp")
            tmp.write_text(json.dumps({"version": VERSION_CACHE, "archivos": vigentes},
                                      ensure_ascii=False, separators=(",", ":")),
                           encoding="utf-8")
            tmp.replace(self.ruta)
        except (OSError, TypeError, ValueError):
            pass                              # sin caché se lee todo, como antes


def _listar_con_firma(carpeta: str) -> List[Tuple[Path, Tuple[int, int]]]:
    """XML de la carpeta (recursivo) con su (mtime_ns, tamaño). Usa scandir: en
    Windows la fecha y el tamaño vienen con el listado, sin tocar cada archivo."""
    import os
    salida, pila = [], [str(carpeta)]
    while pila:
        try:
            with os.scandir(pila.pop()) as it:
                for e in it:
                    try:
                        if e.is_dir(follow_symlinks=False):
                            pila.append(e.path)
                        elif e.name.lower().endswith(".xml"):
                            st = e.stat()
                            salida.append((Path(e.path), (st.st_mtime_ns, st.st_size)))
                    except OSError:
                        continue
        except OSError:
            continue
    salida.sort(key=lambda x: x[0])
    return salida


def leer_carpeta_con_indice(carpeta: str, indice: Indice, progreso=None,
                            forzar: bool = False, cancelado=None
                            ) -> Tuple[List[Renglon], int, int]:
    """Lee la carpeta reusando lo ya leído. Regresa (filas, nuevos, reusados).

    Solo se abren los XML nuevos o modificados (en paralelo); los demás salen de
    la caché de renglones. forzar=True ('Releer todo') ignora la caché.
    """
    from .lector import leer_archivos

    if not Path(carpeta).exists():
        return [], 0, 0
    listado = _listar_con_firma(carpeta)
    cache = CacheFilas(indice.ruta, carpeta)
    en_cache: Dict[str, List[Dict]] = {}
    pendientes: List[Path] = []
    for a, firma in listado:
        previas = None if forzar else cache.filas_de(str(a), firma)
        if previas is None:
            pendientes.append(a)
        else:
            en_cache[str(a)] = previas

    if pendientes or not listado:
        leidas, firmas = leer_archivos(pendientes, progreso=progreso, cancelado=cancelado)
    else:
        leidas, firmas = [], {}
        if progreso:
            progreso(len(listado), len(listado))
    por_archivo: Dict[str, List[Renglon]] = {}
    for r in leidas:
        por_archivo.setdefault(str(r.datos.get("Archivo", "")), []).append(r)

    filas: List[Renglon] = []
    vigentes: Dict[str, list] = {}
    nuevos = reusados = 0
    for a, firma in listado:
        ruta = str(a)
        if ruta in en_cache:
            propias = [Renglon(dict(d)) for d in en_cache[ruta]]
            reusados += len(propias)
            vigentes[ruta] = [firma[0], firma[1], en_cache[ruta]]
        else:
            propias = por_archivo.get(ruta, [])
            nuevos += len(propias)
            if ruta in firmas:                # se pudo abrir (sea o no un CFDI)
                vigentes[ruta] = [firma[0], firma[1], [dict(r.datos) for r in propias]]
        filas.extend(propias)
        if propias:
            r = propias[0]
            indice.datos[ruta] = {
                "archivo": ruta, "mtime": str(firma[0] // 1_000_000_000), "tamano": str(firma[1]),
                "uuid": r.datos.get("UUID", ""), "tipo": r.datos.get("Tipo", ""),
                "fecha": r.datos.get("Fecha", ""), "total": r.datos.get("Total", ""),
            }
    if pendientes or len(vigentes) != len(cache.datos):
        cache.guardar(vigentes)
        indice.guardar()
    return filas, nuevos, reusados


# ---------------------------------------------------------------- export con columnas elegidas
def a_tsv_cols(filas, columnas, con_encabezado=True):
    """TSV solo con las columnas pedidas (para copiar/CSV configurable)."""
    def limpiar(v):
        s = "" if v is None else str(v)
        return s.replace("\t", " ").replace("\r", " ").replace("\n", " ")
    lineas = []
    if con_encabezado:
        lineas.append("\t".join(columnas))
    for r in filas:
        lineas.append("\t".join(limpiar(r.datos.get(c, "")) for c in columnas))
    return "\n".join(lineas)


def a_csv_cols(filas, columnas, destino):
    import csv as _csv
    from pathlib import Path as _P
    d = _P(destino); d.parent.mkdir(parents=True, exist_ok=True)
    with open(d, "w", newline="", encoding="utf-8-sig") as f:
        w = _csv.writer(f)
        w.writerow(columnas)
        for r in filas:
            w.writerow([r.datos.get(c, "") for c in columnas])
    return str(d)


def matriz_a_tsv(cols, filas, con_encabezado=True):
    """TSV a partir de una matriz cruda (resultado de una consulta SQL)."""
    def limpiar(v):
        s = "" if v is None else str(v)
        return s.replace("\t", " ").replace("\r", " ").replace("\n", " ")
    out = []
    if con_encabezado:
        out.append("\t".join(str(c) for c in cols))
    for f in filas:
        out.append("\t".join(limpiar(v) for v in f))
    return "\n".join(out)


def matriz_a_csv(cols, filas, destino):
    import csv as _csv
    from pathlib import Path as _P
    d = _P(destino); d.parent.mkdir(parents=True, exist_ok=True)
    with open(d, "w", newline="", encoding="utf-8-sig") as f:
        w = _csv.writer(f)
        w.writerow(cols)
        for fila in filas:
            w.writerow(fila)
    return str(d)
