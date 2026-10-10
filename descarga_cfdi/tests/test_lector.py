# -*- coding: utf-8 -*-
import csv
from pathlib import Path

import pytest

from sat_descarga.exportar import (Indice, a_csv, a_excel, a_tsv,
                                   leer_carpeta_con_indice)
from sat_descarga.lector import COLUMNAS, leer_carpeta, leer_xml, totales


# ---------------------------------------------------------------- generadores
def cfdi40(uuid="11111111-2222-3333-4444-555555555555", fecha="2026-08-14T10:22:01",
           total="1160.00", subtotal="1000.00", iva="160.00", tipo="I",
           emisor="AAA010101AAA", receptor="HIHR900307P61", serie="A", folio="123",
           conceptos=1, retenciones=""):
    ivas = "".join(
        f'<cfdi:Concepto ClaveProdServ="0101" Cantidad="1" ClaveUnidad="E48" '
        f'Descripcion="Servicio {i+1}" ValorUnitario="1000.00" Importe="1000.00"/>'
        for i in range(conceptos))
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
  xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
  Version="4.0" Serie="{serie}" Folio="{folio}" Fecha="{fecha}"
  SubTotal="{subtotal}" Descuento="0.00" Moneda="MXN" Total="{total}"
  TipoDeComprobante="{tipo}" MetodoPago="PUE" FormaPago="03" LugarExpedicion="06600">
  <cfdi:Emisor Rfc="{emisor}" Nombre="PROVEEDOR SA DE CV" RegimenFiscal="601"/>
  <cfdi:Receptor Rfc="{receptor}" Nombre="RAQUEL HILARIO" UsoCFDI="G03"
    DomicilioFiscalReceptor="06600" RegimenFiscalReceptor="626"/>
  <cfdi:Conceptos>{ivas}</cfdi:Conceptos>
  <cfdi:Impuestos TotalImpuestosTrasladados="{iva}">
    <cfdi:Traslados>
      <cfdi:Traslado Base="{subtotal}" Impuesto="002" TipoFactor="Tasa"
        TasaOCuota="0.160000" Importe="{iva}"/>
    </cfdi:Traslados>
    {retenciones}
  </cfdi:Impuestos>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}"
      FechaTimbrado="{fecha}" RfcProvCertif="SPR190613I52"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode("utf-8")


RETENCIONES = """<cfdi:Retenciones>
      <cfdi:Retencion Impuesto="002" Importe="106.67"/>
      <cfdi:Retencion Impuesto="001" Importe="100.00"/>
    </cfdi:Retenciones>"""


def cfdi33(uuid="99999999-8888-7777-6666-555555555555"):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/3"
  xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
  Version="3.3" Fecha="2019-05-02T09:00:00" SubTotal="500.00" Total="580.00"
  Moneda="MXN" TipoDeComprobante="I" LugarExpedicion="06600">
  <cfdi:Emisor Rfc="BBB020202BBB" Nombre="VIEJO SA" RegimenFiscal="601"/>
  <cfdi:Receptor Rfc="HIHR900307P61" Nombre="RAQUEL" UsoCFDI="G03"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="0101" Cantidad="1" ClaveUnidad="E48"
      Descripcion="Consulta" ValorUnitario="500.00" Importe="500.00"/>
  </cfdi:Conceptos>
  <cfdi:Impuestos TotalImpuestosTrasladados="80.00">
    <cfdi:Traslados>
      <cfdi:Traslado Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000" Importe="80.00"/>
    </cfdi:Traslados>
  </cfdi:Impuestos>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}" FechaTimbrado="2019-05-02T09:05:00"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode("utf-8")


@pytest.fixture
def carpeta(tmp_path):
    """Estructura igual a la que produce la descarga."""
    em = tmp_path / "202608 HIHR" / "emitidos"; em.mkdir(parents=True)
    re_ = tmp_path / "202608 HIHR" / "recibidos"; re_.mkdir(parents=True)
    for i in range(3):
        (em / f"E{i}.xml").write_bytes(cfdi40(uuid=f"1111111{i}-2222-3333-4444-555555555555"))
    for i in range(2):
        (re_ / f"R{i}.xml").write_bytes(cfdi40(uuid=f"2222222{i}-2222-3333-4444-555555555555",
                                               emisor="CCC030303CCC"))
    (em / "leeme.txt").write_text("no soy xml")
    (em / "roto.xml").write_text("<esto no cierra")
    return tmp_path


# ---------------------------------------------------------------- lectura
def test_lee_los_campos_clave(tmp_path):
    a = tmp_path / "x.xml"; a.write_bytes(cfdi40())
    r = leer_xml(a, "Emitido")
    d = r.datos
    assert d["UUID"] == "11111111-2222-3333-4444-555555555555"
    assert d["RFC Emisor"] == "AAA010101AAA" and d["RFC Receptor"] == "HIHR900307P61"
    assert d["SubTotal"] == 1000.0 and d["Total"] == 1160.0
    assert d["IVA Trasladado"] == 160.0
    assert d["Comprobante"] == "Ingreso"
    assert (d["Año"], d["Mes"]) == ("2026", "08")
    assert d["Serie"] == "A" and d["Folio"] == "123"
    assert d["Método Pago"] == "PUE" and d["Uso CFDI"] == "G03"
    assert d["Efecto"] == "Ingreso (+)"
    assert d["Tipo"] == "Emitido"


def test_lee_retenciones(tmp_path):
    a = tmp_path / "x.xml"; a.write_bytes(cfdi40(retenciones=RETENCIONES))
    d = leer_xml(a).datos
    assert d["IVA Retenido"] == 106.67 and d["ISR Retenido"] == 100.0


def test_soporta_version_33(tmp_path):
    a = tmp_path / "v33.xml"; a.write_bytes(cfdi33())
    d = leer_xml(a).datos
    assert d["Total"] == 580.0 and d["IVA Trasladado"] == 80.0


def test_varios_conceptos_resume_descripcion(tmp_path):
    a = tmp_path / "x.xml"; a.write_bytes(cfdi40(conceptos=3))
    d = leer_xml(a).datos
    assert d["Conceptos"] == 3 and "+2 más" in d["Descripción"]


def test_archivo_roto_o_no_cfdi_regresa_none(tmp_path):
    (tmp_path / "a.xml").write_text("<esto no cierra")
    (tmp_path / "b.xml").write_text("<otracosa><x/></otracosa>")
    assert leer_xml(tmp_path / "a.xml") is None
    assert leer_xml(tmp_path / "b.xml") is None
    assert leer_xml(tmp_path / "no_existe.xml") is None


def test_tipo_se_deduce_de_la_carpeta(carpeta):
    filas = leer_carpeta(str(carpeta))
    tipos = sorted(f.datos["Tipo"] for f in filas)
    assert tipos == ["Emitido"] * 3 + ["Recibido"] * 2


def test_lee_carpeta_recursivo_e_ignora_basura(carpeta):
    filas = leer_carpeta(str(carpeta))
    assert len(filas) == 5                      # 5 CFDI; txt y roto fuera


def test_carpeta_inexistente_no_truena(tmp_path):
    assert leer_carpeta(str(tmp_path / "nada")) == []


def test_totales(carpeta):
    t = totales(leer_carpeta(str(carpeta)))
    assert t["Total"] == 5 * 1160.0 and t["IVA Trasladado"] == 5 * 160.0


def test_progreso_se_reporta(carpeta):
    vistos = []
    leer_carpeta(str(carpeta), progreso=lambda i, n: vistos.append((i, n)))
    assert vistos and vistos[-1][0] == vistos[-1][1]


# ---------------------------------------------------------------- portapapeles
def test_tsv_pegable_en_excel(carpeta):
    filas = leer_carpeta(str(carpeta))
    tsv = a_tsv(filas)
    lineas = tsv.split("\n")
    assert lineas[0].split("\t") == COLUMNAS
    assert len(lineas) == 6                     # encabezado + 5
    assert all(len(l.split("\t")) == len(COLUMNAS) for l in lineas)


def test_tsv_sin_encabezado(carpeta):
    assert len(a_tsv(leer_carpeta(str(carpeta)), con_encabezado=False).split("\n")) == 5


def test_tsv_no_rompe_columnas_con_saltos_de_linea(tmp_path):
    x = cfdi40().replace(b'Descripcion="Servicio 1"', b'Descripcion="Linea1&#10;Linea2\tcon tab"')
    (tmp_path / "x.xml").write_bytes(x)
    tsv = a_tsv(leer_carpeta(str(tmp_path)))
    assert len(tsv.split("\n")) == 2
    assert len(tsv.split("\n")[1].split("\t")) == len(COLUMNAS)


# ---------------------------------------------------------------- CSV / Excel
def test_csv_tiene_encabezado_y_filas(carpeta, tmp_path):
    destino = a_csv(leer_carpeta(str(carpeta)), str(tmp_path / "sal" / "cfdi.csv"))
    with open(destino, encoding="utf-8-sig") as f:
        filas = list(csv.reader(f))
    assert filas[0] == COLUMNAS and len(filas) == 6


def test_excel_se_genera_y_se_puede_reabrir(carpeta, tmp_path):
    from openpyxl import load_workbook
    destino = a_excel(leer_carpeta(str(carpeta)), str(tmp_path / "cfdi.xlsx"))
    wb = load_workbook(destino)
    assert wb.sheetnames == ["CFDI", "Resumen"]
    ws = wb["CFDI"]
    assert [c.value for c in ws[1]] == COLUMNAS
    assert ws.max_row == 6
    assert ws.freeze_panes == "A2"
    col_total = COLUMNAS.index("Total") + 1
    assert isinstance(ws.cell(row=2, column=col_total).value, (int, float))


def test_excel_hoja_resumen_agrupa_por_mes_y_tipo(carpeta, tmp_path):
    from openpyxl import load_workbook
    ws = load_workbook(a_excel(leer_carpeta(str(carpeta)), str(tmp_path / "c.xlsx")))["Resumen"]
    filas = [[c.value for c in f] for f in ws.iter_rows(min_row=2)]
    assert len(filas) == 2                                   # Emitido y Recibido
    assert sum(f[2] for f in filas) == 5                     # comprobantes
    assert round(sum(f[7] for f in filas), 2) == 5800.0      # total


def test_excel_vacio_no_truena(tmp_path):
    from openpyxl import load_workbook
    ws = load_workbook(a_excel([], str(tmp_path / "v.xlsx")))["CFDI"]
    assert ws.max_row == 1


# ---------------------------------------------------------------- índice
def test_indice_recuerda_entre_sesiones(carpeta, tmp_path):
    ruta = tmp_path / "indice.csv"
    idx = Indice(str(ruta))
    filas, nuevos, reusados = leer_carpeta_con_indice(str(carpeta), idx)
    assert (nuevos, reusados) == (5, 0) and len(idx) == 5

    idx2 = Indice(str(ruta))                    # nueva sesión
    _, nuevos2, reusados2 = leer_carpeta_con_indice(str(carpeta), idx2)
    assert (nuevos2, reusados2) == (0, 5)


def test_indice_detecta_archivo_nuevo(carpeta, tmp_path):
    ruta = tmp_path / "i.csv"
    idx = Indice(str(ruta))
    leer_carpeta_con_indice(str(carpeta), idx)
    (carpeta / "202608 HIHR" / "emitidos" / "NUEVO.xml").write_bytes(
        cfdi40(uuid="33333333-2222-3333-4444-555555555555"))
    _, nuevos, reusados = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert (nuevos, reusados) == (1, 5)


def test_indice_detecta_archivo_modificado(carpeta, tmp_path):
    import os, time
    ruta = tmp_path / "i.csv"
    idx = Indice(str(ruta))
    leer_carpeta_con_indice(str(carpeta), idx)
    objetivo = carpeta / "202608 HIHR" / "emitidos" / "E0.xml"
    objetivo.write_bytes(cfdi40(uuid="44444444-2222-3333-4444-555555555555", total="2000.00"))
    os.utime(objetivo, (time.time() + 10, time.time() + 10))
    _, nuevos, _ = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert nuevos == 1


def test_indice_forzar_reprocesa_todo(carpeta, tmp_path):
    ruta = tmp_path / "i.csv"
    leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    _, nuevos, reusados = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)), forzar=True)
    assert (nuevos, reusados) == (5, 0)


def test_indice_olvida_borrados(carpeta, tmp_path):
    ruta = tmp_path / "i.csv"
    idx = Indice(str(ruta))
    leer_carpeta_con_indice(str(carpeta), idx)
    (carpeta / "202608 HIHR" / "emitidos" / "E0.xml").unlink()
    assert idx.olvidar_inexistentes() == 1 and len(idx) == 4


def test_indice_corrupto_no_truena(tmp_path):
    ruta = tmp_path / "i.csv"; ruta.write_text("basura,,,\nsin sentido")
    assert len(Indice(str(ruta))) == 0


# ---- generadores para pagos y relacionados (usados por varios tests) ----
def cfdi_pago(uuid="AAAA1111-2222-3333-4444-555555555555",
              docto_uuid="BBBB1111-2222-3333-4444-555555555555",
              monto="2320.00", imp_pagado="2320.00", parcialidad="1",
              saldo_ant="2320.00", saldo_insoluto="0.00", fecha="2026-08-12T10:28:04"):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
  xmlns:pago20="http://www.sat.gob.mx/Pagos20"
  xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
  Version="4.0" Serie="P" Folio="900" Fecha="{fecha}" SubTotal="0" Total="0"
  Moneda="XXX" TipoDeComprobante="P" LugarExpedicion="06600">
  <cfdi:Emisor Rfc="IMO011025RIA" Nombre="INSIGNIA MOTORS" RegimenFiscal="601"/>
  <cfdi:Receptor Rfc="FODA920130KD4" Nombre="ADRIANA FONSECA" UsoCFDI="CP01"
    DomicilioFiscalReceptor="06600" RegimenFiscalReceptor="626"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="84111506" Cantidad="1" ClaveUnidad="ACT"
      Descripcion="Pago" ValorUnitario="0" Importe="0" ObjetoImp="01"/>
  </cfdi:Conceptos>
  <cfdi:Complemento>
    <pago20:Pagos Version="2.0">
      <pago20:Totales MontoTotalPagos="{monto}"/>
      <pago20:Pago FechaPago="{fecha}" FormaDePagoP="03" MonedaP="MXN" TipoCambioP="1"
        Monto="{monto}">
        <pago20:DoctoRelacionado IdDocumento="{docto_uuid}" MonedaDR="MXN"
          NumParcialidad="{parcialidad}" ImpSaldoAnt="{saldo_ant}"
          ImpPagado="{imp_pagado}" ImpSaldoInsoluto="{saldo_insoluto}"
          ObjetoImpDR="02">
          <pago20:ImpuestosDR>
            <pago20:TrasladosDR>
              <pago20:TrasladoDR BaseDR="2000.00" ImpuestoDR="002" TipoFactorDR="Tasa"
                TasaOCuotaDR="0.160000" ImporteDR="320.00"/>
            </pago20:TrasladosDR>
          </pago20:ImpuestosDR>
        </pago20:DoctoRelacionado>
      </pago20:Pago>
    </pago20:Pagos>
  </cfdi:Complemento>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}" FechaTimbrado="{fecha}"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode("utf-8")


def cfdi_con_relacionado(uuid="CCCC1111-2222-3333-4444-555555555555",
                          rel_uuid="DDDD1111-2222-3333-4444-555555555555",
                          tipo_rel="01"):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
  xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
  Version="4.0" Serie="E" Folio="5" Fecha="2026-08-20T10:00:00"
  SubTotal="500.00" Total="580.00" Moneda="MXN" TipoDeComprobante="E" LugarExpedicion="06600">
  <cfdi:CfdiRelacionados TipoRelacion="{tipo_rel}">
    <cfdi:CfdiRelacionado UUID="{rel_uuid}"/>
  </cfdi:CfdiRelacionados>
  <cfdi:Emisor Rfc="AAA010101AAA" Nombre="PROVEEDOR" RegimenFiscal="601"/>
  <cfdi:Receptor Rfc="HIHR900307P61" Nombre="RAQUEL" UsoCFDI="G02"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="0101" Cantidad="1" ClaveUnidad="E48"
      Descripcion="Devolución" ValorUnitario="500.00" Importe="500.00"/>
  </cfdi:Conceptos>
  <cfdi:Impuestos TotalImpuestosTrasladados="80.00">
    <cfdi:Traslados>
      <cfdi:Traslado Base="500" Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000" Importe="80.00"/>
    </cfdi:Traslados>
  </cfdi:Impuestos>
  <cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}" FechaTimbrado="2026-08-20T10:00:00"/>
  </cfdi:Complemento>
</cfdi:Comprobante>""".encode("utf-8")


# ---- pagos y relacionados en el renglón ----
def test_pago_extrae_monto_real(tmp_path):
    from sat_descarga.lector import leer_xml
    x = tmp_path / "p.xml"; x.write_bytes(cfdi_pago(monto="2320.00"))
    r = leer_xml(x, "Recibido")
    assert r.datos["Total"] == 0.0                        # el comprobante es 0
    assert r.datos["Monto Total Pago"] == 2320.0         # pero el pago real sí sale
    assert r.datos["Imp Pagado"] == 2320.0               # importe del doc pagado
    assert r.datos["Comprobante"] == "Pago"


def test_relacionado_en_renglon(tmp_path):
    from sat_descarga.lector import leer_xml
    x = tmp_path / "r.xml"; x.write_bytes(cfdi_con_relacionado(rel_uuid="DDDD1111-2222-3333-4444-555555555555"))
    r = leer_xml(x, "Recibido")
    assert "DDDD1111" in r.datos["Relacionados"]


def test_detalle_pago_completo(tmp_path):
    from sat_descarga.lector import leer_detalle
    x = tmp_path / "p.xml"; x.write_bytes(cfdi_pago(monto="5000.00", imp_pagado="5000.00"))
    d = leer_detalle(x)
    assert d["es_pago"] is True
    assert d["pago_monto"] == 5000.0
    assert len(d["pago_docs"]) == 1
    assert d["pago_docs"][0]["pagado"] == 5000.0
    assert d["pago_docs"][0]["uuid"]


def test_detalle_relacionados(tmp_path):
    from sat_descarga.lector import leer_detalle
    x = tmp_path / "r.xml"; x.write_bytes(cfdi_con_relacionado(tipo_rel="01"))
    d = leer_detalle(x)
    assert len(d["relacionados"]) == 1
    assert d["relacionados"][0]["tipo_desc"] == "Nota de crédito"


def test_cfdi_normal_sin_pago_ni_relacionados(tmp_path):
    from sat_descarga.lector import leer_xml, leer_detalle
    x = tmp_path / "n.xml"; x.write_bytes(cfdi40())
    r = leer_xml(x, "Emitido")
    assert r.datos["Monto Total Pago"] == 0.0 and r.datos["Relacionados"] == ""
    d = leer_detalle(x)
    assert d["pago_docs"] == [] and d["relacionados"] == []


def test_pago_extrae_iva_y_operacion(tmp_path):
    from sat_descarga.lector import leer_xml, leer_detalle
    x = tmp_path / "p.xml"; x.write_bytes(cfdi_pago())
    r = leer_xml(x, "Recibido")
    assert r.datos["IVA Doc"] == 320.0
    assert r.datos["Doc Pagado UUID"] == "BBBB1111-2222-3333-4444-555555555555"
    assert r.datos["Forma Pago P"] == "03"
    d = leer_detalle(x)
    assert d["pago_iva"] == 320.0
    assert d["pago_docs"][0]["iva"] == 320.0


def test_volcado_xml_saca_todo_menos_sellos(tmp_path):
    from sat_descarga.lector import volcado_xml
    x = tmp_path / "p.xml"; x.write_bytes(cfdi_pago())
    filas = volcado_xml(x)
    assert len(filas) > 20
    # tiene los datos clave
    campos = {c for _, c, _ in filas}
    assert "Version" in campos and "Total" in campos and "UUID" in campos
    assert "IdDocumento" in campos          # del complemento de pago
    # NO tiene sellos ni cadenas
    assert not any(c.lower() in ("sello", "sellocfd", "sellosat", "certificado",
                                 "cadenaoriginal") for _, c, _ in filas)


def test_volcado_uuid_si_aparece(tmp_path):
    from sat_descarga.lector import volcado_xml
    x = tmp_path / "f.xml"; x.write_bytes(cfdi40())
    filas = volcado_xml(x)
    assert any(c == "UUID" for _, c, _ in filas)


def test_volcado_no_cfdi_vacio(tmp_path):
    from sat_descarga.lector import volcado_xml
    (tmp_path / "b.xml").write_text("<x/>")
    assert isinstance(volcado_xml(tmp_path / "b.xml"), list)


# ---- XML real de Rick: REP con 2 documentos pagados ----
def test_rep_real_expande_una_fila_por_documento():
    from sat_descarga.lector import leer_xml_filas
    ruta = Path(__file__).parent / "xml_reales" / "pago_133k.xml"
    filas = leer_xml_filas(ruta, "Recibido")
    assert len(filas) == 2                            # 2 documentos pagados

    d0, d1 = filas[0].datos, filas[1].datos
    # ambos comparten el monto total del pago
    assert d0["Monto Total Pago"] == 133889.04
    assert d0["Num Operación"] == "17"
    assert d0["Régimen Emisor"] == "601"
    # cada fila su propio documento
    docs = {f.datos["Doc Folio"]: f.datos for f in filas}
    assert docs["15761"]["Imp Pagado"] == 9500.0
    assert docs["15761"]["IVA Doc"] == 703.70
    assert docs["4237"]["Imp Pagado"] == 124389.04
    assert docs["4237"]["IVA Doc"] == 9214.0
    assert docs["4237"]["Saldo Insoluto"] == 45000.0
    assert docs["15761"]["Tasa Doc"] == "0.080000"


def test_rep_real_columnas_completas():
    from sat_descarga.lector import leer_xml_filas, COLUMNAS
    ruta = Path(__file__).parent / "xml_reales" / "pago_133k.xml"
    f = leer_xml_filas(ruta, "Recibido")[0]
    # todas las columnas presentes
    for c in COLUMNAS:
        assert c in f.datos, f"falta columna {c}"
    assert f.datos["Es Pago"] == "Sí"
    assert f.datos["Uso CFDI"] == "CP01"
    assert f.datos["Domicilio Receptor"] == "32546"


# ---------------------------------------------------------------- caché de renglones
def _contar_aperturas(monkeypatch):
    """Cuenta cuántos XML se abren de verdad."""
    import sat_descarga.lector as lector
    abiertos = []
    real = lector._bytes_de
    monkeypatch.setattr(lector, "_bytes_de", lambda a: abiertos.append(a) or real(a))
    return abiertos


def test_cache_no_vuelve_a_abrir_los_xml(carpeta, tmp_path, monkeypatch):
    ruta = tmp_path / "i.csv"
    primeras, _, _ = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    abiertos = _contar_aperturas(monkeypatch)
    filas, nuevos, reusados = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert abiertos == [] and (nuevos, reusados) == (0, 5)
    assert [r.datos for r in filas] == [r.datos for r in primeras]      # mismos datos, mismo orden


def test_cache_solo_abre_lo_nuevo_o_modificado(carpeta, tmp_path, monkeypatch):
    import os, time
    ruta = tmp_path / "i.csv"
    leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    em = carpeta / "202608 HIHR" / "emitidos"
    (em / "NUEVO.xml").write_bytes(cfdi40(uuid="33333333-2222-3333-4444-555555555555"))
    (em / "E0.xml").write_bytes(cfdi40(uuid="44444444-2222-3333-4444-555555555555", total="2000.00"))
    os.utime(em / "E0.xml", (time.time() + 10, time.time() + 10))
    abiertos = _contar_aperturas(monkeypatch)
    filas, nuevos, reusados = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert sorted(a.name for a in abiertos) == ["E0.xml", "NUEVO.xml"]
    assert (nuevos, reusados) == (2, 4)
    assert any(r.datos["Total"] == 2000.0 for r in filas)


def test_cache_releer_todo_la_ignora_y_borrados_salen(carpeta, tmp_path, monkeypatch):
    ruta = tmp_path / "i.csv"
    leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    abiertos = _contar_aperturas(monkeypatch)
    _, nuevos, _ = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)), forzar=True)
    assert nuevos == 5 and len(abiertos) == 6            # 5 CFDI + roto.xml
    (carpeta / "202608 HIHR" / "emitidos" / "E1.xml").unlink()
    filas, _, _ = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert len(filas) == 4


def test_cache_danada_se_ignora(carpeta, tmp_path):
    ruta = tmp_path / "i.csv"
    leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    for f in (tmp_path / "cache_filas").glob("*.json"):
        f.write_text("{roto", encoding="utf-8")
    filas, nuevos, _ = leer_carpeta_con_indice(str(carpeta), Indice(str(ruta)))
    assert len(filas) == 5 and nuevos == 5
