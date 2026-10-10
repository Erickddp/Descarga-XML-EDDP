# -*- coding: utf-8 -*-
"""Pruebas de la interfaz gráfica: se levanta Tk de verdad (pantalla virtual)
y se manipulan los widgets como lo haría el usuario. No se abre ningún diálogo
real: se interceptan filedialog/messagebox.
"""
import json
from datetime import date
from pathlib import Path

import pytest

tk = pytest.importorskip("tkinter")


@pytest.fixture(scope="module")
def _tk_disponible():
    try:
        r = tk.Tk(); r.withdraw(); r.destroy()
    except tk.TclError as e:
        pytest.skip(f"sin display: {e}")


@pytest.fixture
def app(tmp_path, monkeypatch, _tk_disponible):
    monkeypatch.chdir(tmp_path)
    import app as modulo
    a = modulo.App(ruta_config=str(tmp_path / "config.json"))
    a.withdraw()
    yield a
    try:
        aid = getattr(a, "_after_cola", None)
        if aid:
            a.after_cancel(aid)
    except tk.TclError:
        pass
    try:
        a.destroy()
    except tk.TclError:
        pass


def _archivos_efirma(tmp_path, rfc="AAA010101AAA"):
    cer = tmp_path / f"{rfc}.cer"; cer.write_bytes(b"x")
    key = tmp_path / f"{rfc}.key"; key.write_bytes(b"x")
    return str(cer), str(key)


# ------------------------------------------------------------ arranque
def test_arranca_sin_config_previa(app):
    assert app.cfg.perfiles == []
    assert app.tabla.get_children() == ()


def test_fechas_default_son_del_anio_en_curso(app):
    assert app.v_desde.get() == f"{date.today().year}-01-01"
    assert app.v_hasta.get() == date.today().isoformat()


# ------------------------------------------------------------ diálogo de alta
def test_dialogo_selector_de_archivo_llena_ruta_y_detecta_rfc(app, tmp_path, monkeypatch):
    """El usuario da clic en 'Buscar…', escoge el .key donde se le dé la gana,
    y la app rellena la ruta, el RFC y encuentra el .cer hermano."""
    import app as modulo
    cer, key = _archivos_efirma(tmp_path, "EDDP800101XY1")
    monkeypatch.setattr(modulo.filedialog, "askopenfilename", lambda **k: key)

    d = modulo.DialogoPerfil(app)
    d._buscar_key()
    assert d.v_key.get() == key
    assert d.v_rfc.get() == "EDDP800101XY1"     # sacado del nombre del archivo
    assert d.v_cer.get() == cer                 # encontró el .cer junto a la llave
    d.destroy()


def test_dialogo_rechaza_datos_incompletos(app, monkeypatch):
    import app as modulo
    errores = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: errores.append(a))
    d = modulo.DialogoPerfil(app)
    d.v_rfc.set("CORTO")
    d._guardar()
    assert d.resultado is None and errores      # no guardó y avisó
    d.destroy()


def test_dialogo_guarda_perfil_valido(app, tmp_path):
    import app as modulo
    cer, key = _archivos_efirma(tmp_path)
    d = modulo.DialogoPerfil(app)
    d.v_rfc.set("aaa010101aaa"); d.v_cer.set(cer); d.v_key.set(key)
    d.v_pass.set("secreta"); d.v_nombre.set("Cliente A")
    d._guardar()
    assert d.resultado.rfc == "AAA010101AAA" and d.resultado.nombre == "Cliente A"


def test_ver_contrasena_alterna(app):
    import app as modulo
    d = modulo.DialogoPerfil(app)
    assert d.e_pass.cget("show") == "•"
    d.v_ver.set(True); d._toggle()
    assert d.e_pass.cget("show") == ""
    d.destroy()


# ------------------------------------------------------------ lista de RFCs
def _alta(app, tmp_path, rfc, nombre):
    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, rfc)
    app.cfg.agregar(Perfil(rfc=rfc, cer_path=cer, key_path=key, key_password="p", nombre=nombre))
    app._guardar_cfg(); app._refrescar_lista()


def test_varios_rfcs_se_listan_y_persisten(app, tmp_path):
    _alta(app, tmp_path, "AAA010101AAA", "Cliente A")
    _alta(app, tmp_path, "BBB020202BBB", "Cliente B")
    _alta(app, tmp_path, "CCC030303CCC", "Cliente C")
    _alta(app, tmp_path, "DDD040404DDD", "Cliente D")
    assert len(app.tabla.get_children()) == 4

    guardado = json.loads((tmp_path / "config.json").read_text(encoding="utf-8"))
    assert len(guardado["perfiles"]) == 4
    assert guardado["perfiles"][0]["nombre"] == "Cliente A"

    import app as modulo
    otra = modulo.App(ruta_config=str(tmp_path / "config.json")); otra.withdraw()
    assert len(otra.cfg.perfiles) == 4
    otra.destroy()


def test_marcar_todos_y_ninguno(app, tmp_path):
    _alta(app, tmp_path, "AAA010101AAA", "A")
    _alta(app, tmp_path, "BBB020202BBB", "B")
    app._marcar(False)
    assert app.cfg.activos() == []
    app._marcar(True)
    assert len(app.cfg.activos()) == 2
    fila = app.cfg.por_rfc("AAA010101AAA").id
    assert app.tabla.item(fila)["values"][0] == "☑"


def test_quitar_rfc_seleccionado(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "AAA010101AAA", "A")
    _alta(app, tmp_path, "BBB020202BBB", "B")
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: True)
    app.tabla.selection_set(app.cfg.por_rfc("AAA010101AAA").id)
    app._quitar()
    assert [p.rfc for p in app.cfg.perfiles] == ["BBB020202BBB"]
    assert len(app.tabla.get_children()) == 1


def test_estado_marca_perfiles_incompletos(app, tmp_path):
    from sat_descarga.config import Perfil
    app.cfg.agregar(Perfil(rfc="AAA010101AAA", cer_path="/no/existe.cer",
                           key_path="/no/existe.key", key_password="", nombre="Roto"))
    app._refrescar_lista()
    fila = app.cfg.por_rfc("AAA010101AAA").id
    assert app.tabla.item(fila)["values"][5] == "revisar datos"


# ------------------------------------------------------------ validaciones al descargar
def test_no_descarga_si_fechas_invalidas(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "AAA010101AAA", "A")
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: avisos.append(a))
    app.v_desde.set("01/01/2026")
    app._descargar()
    assert avisos and app.hilo is None


def test_no_descarga_si_desde_mayor_que_hasta(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "AAA010101AAA", "A")
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: avisos.append(a))
    app.v_desde.set("2026-09-05"); app.v_hasta.set("2026-01-01")
    app._descargar()
    assert avisos and app.hilo is None


def test_avisa_si_no_hay_rfcs_marcados(app, monkeypatch):
    import app as modulo
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showwarning", lambda *a, **k: avisos.append(a))
    app._descargar()
    assert avisos and "RFCs marcados" in avisos[0][1]


def test_descarga_real_end_to_end_en_la_ui(app, tmp_path, monkeypatch):
    """Aprieta el botón Descargar con un SAT simulado y verifica que los XML
    aparecen en disco y el botón se rehabilita."""
    import app as modulo
    from conftest import ClienteFalso, cfdi_xml, zip_b64

    _alta(app, tmp_path, "AAA010101AAA", "A")
    _alta(app, tmp_path, "BBB020202BBB", "B")
    app.v_desde.set("2026-01-01"); app.v_hasta.set("2026-01-31")
    app.v_tipo.set("emitidos")
    app.v_salida.set(str(tmp_path / "salida"))
    app.cfg.opciones.poll_intervalo_seg = 0

    paq = {"P1": zip_b64({"a.xml": cfdi_xml("2026-01-10T00:00:00")})}
    original = modulo.ejecutar_varios
    monkeypatch.setattr(modulo, "ejecutar_varios", lambda *a, **k: original(
        *a, **{**k, "dormir": lambda s: None,
               "fabrica_cliente": lambda p: ClienteFalso(paq, verificaciones_en_proceso=0)}))
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    monkeypatch.setattr(modulo.messagebox, "showwarning", lambda *a, **k: None)

    app._descargar()
    assert app.hilo is not None
    app.hilo.join(timeout=30)
    assert not app.hilo.is_alive(), "el hilo de descarga no terminó"

    # drenar la cola hasta que la UI procese el mensaje final (determinista)
    import time as _t
    limite = _t.time() + 10
    while "Terminado" not in app.v_estado.get() and _t.time() < limite:
        app._vaciar_cola()
        app.update()
        _t.sleep(0.02)

    salida = Path(tmp_path / "salida")
    assert list((salida / "202601 AAA0" / "emitidos").glob("*.xml"))
    assert list((salida / "202601 BBB0" / "emitidos").glob("*.xml"))
    assert str(app.btn_desc.cget("state")) == "normal"
    assert "Terminado" in app.v_estado.get()


def test_atajos_de_fecha(app):
    app._anio_pasado()
    a = date.today().year - 1
    assert app.v_desde.get() == f"{a}-01-01" and app.v_hasta.get() == f"{a}-12-31"
    app._mes_pasado()
    assert app.v_desde.get().endswith("-01")
    assert app.v_desde.get()[:7] == app.v_hasta.get()[:7]


def test_carpeta_salida_se_guarda_en_config(app, tmp_path):
    destino = str(tmp_path / "mis_xml")
    app.v_salida.set(destino)
    app._guardar_cfg()
    guardado = json.loads((tmp_path / "config.json").read_text(encoding="utf-8"))
    assert guardado["opciones"]["carpeta_salida"] == destino


def test_log_de_la_app_llega_al_recuadro(app):
    app.log.info("mensaje de prueba 12345")
    app._vaciar_cola()
    app.update()
    assert "mensaje de prueba 12345" in app.txt.get("1.0", "end")


class _DialogoFalso:
    """Doble del diálogo ya cerrado: solo entrega el perfil capturado."""
    def __init__(self, perfil):
        self.resultado = perfil


def _simular_alta_por_dialogo(app, monkeypatch, perfil):
    import app as modulo
    monkeypatch.setattr(modulo, "DialogoPerfil", lambda *a, **k: _DialogoFalso(perfil))
    monkeypatch.setattr(app, "wait_window", lambda *a, **k: None)
    app._agregar()


# ============================================================ regresión del bug real
def test_agregar_segundo_rfc_no_borra_el_primero_en_la_ui(app, tmp_path, monkeypatch):
    """El escenario exacto: se dio de alta Raquel, luego Adriana, y Raquel
    desapareció de la tabla."""
    import app as modulo
    _alta(app, tmp_path, "HIHR900307P61", "RAQUEL HILARIO H")

    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, "FODA920130KD4")
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: True)
    _simular_alta_por_dialogo(app, monkeypatch, Perfil(
        rfc="FODA920130KD4", cer_path=cer, key_path=key,
        key_password="p", nombre="ADRIANA FONSECA"))

    nombres = sorted(p.nombre for p in app.cfg.perfiles)
    assert nombres == ["ADRIANA FONSECA", "RAQUEL HILARIO H"]
    assert len(app.tabla.get_children()) == 2


def test_avisa_pero_conserva_ambos_si_el_rfc_se_repite(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "HIHR900307P61", "RAQUEL")
    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, "otro")
    preguntas = []
    monkeypatch.setattr(modulo.messagebox, "askyesno",
                        lambda *a, **k: (preguntas.append(a), True)[1])
    _simular_alta_por_dialogo(app, monkeypatch, Perfil(
        rfc="HIHR900307P61", cer_path=cer, key_path=key, key_password="p", nombre="ADRIANA"))
    assert preguntas, "debió avisar que el RFC ya existía"
    assert len(app.cfg.perfiles) == 2


def test_si_dice_que_no_al_duplicado_no_agrega_nada(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "HIHR900307P61", "RAQUEL")
    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, "otro2")
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: False)
    _simular_alta_por_dialogo(app, monkeypatch, Perfil(
        rfc="HIHR900307P61", cer_path=cer, key_path=key, key_password="p", nombre="ADRIANA"))
    assert len(app.cfg.perfiles) == 1


def test_el_certificado_corrige_un_rfc_mal_escrito(app, tmp_path, monkeypatch):
    """Si el usuario deja pegado el RFC anterior, el .cer manda."""
    import app as modulo
    from conftest import openssl
    openssl("req", "-x509", "-newkey", "rsa:2048", "-nodes",
            "-keyout", tmp_path / "k.pem", "-out", tmp_path / "c.pem",
            "-days", "1", "-subj", "/CN=ADRIANA FONSECA/serialNumber=FODA920130KD4")
    cer = tmp_path / "cualquier_nombre.cer"
    openssl("x509", "-in", tmp_path / "c.pem", "-outform", "DER", "-out", cer)

    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showwarning", lambda *a, **k: avisos.append(a))
    d = modulo.DialogoPerfil(app)
    d.v_rfc.set("HIHR900307P61")          # RFC equivocado pegado
    d._leer_certificado(str(cer))
    assert d.v_rfc.get() == "FODA920130KD4"
    assert avisos, "debió avisar que el certificado era de otro RFC"
    d.destroy()


def test_las_contrasenas_y_rutas_sobreviven_al_reinicio(app, tmp_path):
    """No debe volver a capturar e.firma cada vez que abre."""
    import app as modulo
    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, "HIHR900307P61")
    app.cfg.agregar(Perfil(rfc="HIHR900307P61", cer_path=cer, key_path=key,
                           key_password="MiClaveSecreta", nombre="RAQUEL"))
    app._guardar_cfg()

    otra = modulo.App(ruta_config=str(tmp_path / "config.json")); otra.withdraw()
    p = otra.cfg.por_rfc("HIHR900307P61")
    assert p.key_password == "MiClaveSecreta"
    assert p.cer_path == cer and p.key_path == key
    assert p.valido()
    otra.destroy()


# ============================================================ pestaña Datos / Excel
def _poblar_xml(carpeta_base, n_em=3, n_re=2):
    """Crea XML como los que deja la descarga."""
    from test_lector import cfdi40
    em = Path(carpeta_base) / "202608 HIHR" / "emitidos"; em.mkdir(parents=True, exist_ok=True)
    rc = Path(carpeta_base) / "202608 HIHR" / "recibidos"; rc.mkdir(parents=True, exist_ok=True)
    for i in range(n_em):
        (em / f"E{i}.xml").write_bytes(cfdi40(uuid=f"1111111{i}-2222-3333-4444-555555555555"))
    for i in range(n_re):
        (rc / f"R{i}.xml").write_bytes(cfdi40(uuid=f"2222222{i}-2222-3333-4444-555555555555",
                                             emisor="CCC030303CCC"))
    return str(carpeta_base)


def test_pestana_datos_existe(app):
    assert "Datos" in " ".join(app.nb.tab(i, "text") for i in range(app.nb.index("end")))


def test_leer_xml_llena_la_tabla(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    assert len(app.filas) == 5
    assert len(app.tabla_datos.get_children()) == 5
    assert "5" in app.v_totales.get()
    assert app._chips["cfdi"].cget("text") == "5"


def test_totales_se_muestran(app, tmp_path):
    """3 emitidos (I) y 2 recibidos (G) de $1,160 con $160 de IVA cada uno."""
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    assert app._chips["cfdi"].cget("text") == "5"
    assert "3,480.00" in app._chips["ingresos"].cget("text")
    assert "2,320.00" in app._chips["gastos"].cget("text")
    assert "480.00" in app._chips["iva_i"].cget("text")
    assert "320.00" in app._chips["iva_g"].cget("text")
    assert "3 facturas emitidas" in app._chips_sub["ingresos"].cget("text")


def test_filtro_por_tipo(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.v_f_tipo.set("Recibido"); app._aplicar_filtros()
    assert len(app.filas_vista) == 2
    assert all(r.datos["Tipo"] == "Recibido" for r in app.filas_vista)


def test_filtro_por_texto_busca_en_varios_campos(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.v_buscar.set("ccc030303"); app._aplicar_filtros()
    assert len(app.filas_vista) == 2                  # solo los del emisor CCC
    app.v_buscar.set("noexiste"); app._aplicar_filtros()
    assert app.filas_vista == []


def test_limpiar_filtros_restaura(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.v_f_tipo.set("Emitido"); app.v_buscar.set("zzz"); app._aplicar_filtros()
    app._limpiar_filtros()
    assert len(app.filas_vista) == 5


def test_combo_de_meses_se_llena_solo(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    assert "2026-08" in app.cb_mes.cget("values")


def test_ordenar_por_columna_alterna(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app._ordenar_por("UUID")
    asc = [r.datos["UUID"] for r in app.filas_vista]
    app._ordenar_por("UUID")
    assert [r.datos["UUID"] for r in app.filas_vista] == list(reversed(asc))


def test_copiar_tabla_deja_tsv_en_portapapeles(app, tmp_path):
    from sat_descarga.lector import COLUMNAS_COPIADO
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app._copiar_tabla()
    texto = app.clipboard_get()
    lineas = texto.split("\n")
    assert lineas[0].split("\t") == COLUMNAS_COPIADO
    assert len(lineas) == 6
    assert "Copiado: 5" in app.v_totales.get()


def test_copiar_respeta_el_filtro(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.v_f_tipo.set("Recibido"); app._aplicar_filtros()
    app._copiar_tabla()
    assert len(app.clipboard_get().split("\n")) == 3          # encabezado + 2


def test_copiar_seleccion(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    hijos = app.tabla_datos.get_children()
    app.tabla_datos.selection_set([hijos[0], hijos[2]])
    app._copiar_seleccion()
    assert len(app.clipboard_get().split("\n")) == 3


def test_copiar_sin_datos_avisa(app, monkeypatch):
    import app as modulo
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: avisos.append(a))
    app._copiar_tabla()
    assert avisos


def test_exportar_excel_desde_la_ui(app, tmp_path, monkeypatch):
    import app as modulo
    from openpyxl import load_workbook
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    destino = str(tmp_path / "salida.xlsx")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_excel()
    wb = load_workbook(destino)
    assert wb["CFDI"].max_row == 6 and "Resumen" in wb.sheetnames


def test_exportar_excel_exporta_lo_filtrado(app, tmp_path, monkeypatch):
    import app as modulo
    from openpyxl import load_workbook
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.v_f_tipo.set("Emitido"); app._aplicar_filtros()
    destino = str(tmp_path / "filtrado.xlsx")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_excel()
    assert load_workbook(destino)["CFDI"].max_row == 4        # encabezado + 3


def test_cancelar_el_guardado_no_truena(app, tmp_path, monkeypatch):
    import app as modulo
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: "")
    app._exportar_excel()
    app._exportar_csv()


def test_exportar_csv_desde_la_ui(app, tmp_path, monkeypatch):
    import csv as _csv
    import app as modulo
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    destino = str(tmp_path / "s.csv")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_csv()
    with open(destino, encoding="utf-8-sig") as f:
        assert len(list(_csv.reader(f))) == 6


def test_carpeta_inexistente_avisa(app, monkeypatch):
    import app as modulo
    errores = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: errores.append(a))
    app.v_carpeta_datos.set("/ruta/que/no/existe")
    app._leer_datos(esperar=True)
    assert errores


def test_carpeta_vacia_avisa_pero_no_truena(app, tmp_path, monkeypatch):
    import app as modulo
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: avisos.append(a))
    vacia = tmp_path / "vacia"; vacia.mkdir()
    app.v_carpeta_datos.set(str(vacia))
    app._leer_datos(esperar=True)
    assert avisos and app.filas == []


def test_segunda_lectura_reusa_el_indice(app, tmp_path):
    carpeta = _poblar_xml(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta)
    app._leer_datos(esperar=True)
    app._leer_datos(esperar=True)                                  # no debe duplicar ni fallar
    assert len(app.filas) == 5
    assert (Path(app.cfg.ruta).parent / "indice_xml.csv").exists()


# ============================================================ tema y sesión en la UI
def test_app_recuerda_carpeta_de_datos(app, tmp_path, monkeypatch):
    import app as modulo
    carpeta = _poblar_xml(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta)
    app._leer_datos(esperar=True)
    # relee config desde disco y confirma que quedó
    from sat_descarga.config import Config
    c = Config.desde_archivo(app.cfg.ruta)
    assert c.opciones.ultima_carpeta_datos == carpeta


def test_app_recuerda_fechas_y_tipo(app, tmp_path):
    from sat_descarga.config import Config
    app.v_desde.set("2025-03-01"); app.v_hasta.set("2025-03-31"); app.v_tipo.set("recibidos")
    app._guardar_cfg()
    c = Config.desde_archivo(app.cfg.ruta)
    assert c.opciones.ultimo_desde == "2025-03-01"
    assert c.opciones.ultimo_tipo == "recibidos"


def test_al_reabrir_usa_la_ultima_carpeta(app, tmp_path):
    import app as modulo
    carpeta = _poblar_xml(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta); app._leer_datos(esperar=True)
    otra = modulo.App(ruta_config=app.cfg.ruta); otra.withdraw()
    assert otra.v_carpeta_datos.get() == carpeta
    otra.destroy()


def test_alternar_tema_guarda_preferencia(app, monkeypatch):
    import app as modulo
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    inicial = app.tema.nombre
    app._alternar_tema()
    from sat_descarga.config import Config
    c = Config.desde_archivo(app.cfg.ruta)
    assert c.opciones.tema != inicial


def test_chips_existen_y_arrancan_vacios(app):
    assert set(app._chips) == {"cfdi", "ingresos", "gastos", "iva_i", "iva_g", "saldo"}
    assert app._chips["cfdi"].cget("text") == "—"


def test_encabezado_muestra_conteo_de_rfcs(app, tmp_path):
    _alta(app, tmp_path, "AAA010101AAA", "A")
    _alta(app, tmp_path, "BBB020202BBB", "B")
    assert "2 RFC" in app.lbl_ident.cget("text")


# ============================================================ layout / scroll
def test_controles_de_descarga_existen(app):
    """Regresión: los botones y campos de descarga deben existir aunque la
    ventana sea corta (van dentro de un área con scroll)."""
    assert app.btn_desc.cget("text").strip().endswith("Descargar XML")
    assert app.btn_stop is not None
    assert app.v_salida.get()          # campo carpeta de salida
    assert app.v_desde.get() and app.v_hasta.get()


def test_existe_area_con_scroll_en_descargar(app):
    import tkinter as tk
    def hay_canvas(w):
        if isinstance(w, tk.Canvas):
            return True
        return any(hay_canvas(c) for c in w.winfo_children())
    assert hay_canvas(app), "la pestaña Descargar debe tener un canvas con scroll"


def test_boton_descargar_es_accionable(app, tmp_path, monkeypatch):
    import app as modulo
    _alta(app, tmp_path, "AAA010101AAA", "A")
    # sin fechas válidas debe avisar, probando que el botón está cableado
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: avisos.append(a))
    app.v_desde.set("mal")
    app._descargar()
    assert avisos


# ============================================================ dashboard / SQL / columnas
def test_subpestanas_existen(app):
    textos = [app.sub.tab(i, "text") for i in range(app.sub.index("end"))]
    assert any("Tabla" in t for t in textos)
    assert any("Panel" in t for t in textos)
    assert any("SQL" in t for t in textos)


def test_leer_construye_motor_sql(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    assert app.motor is not None
    cols, filas = app.motor.consultar("SELECT COUNT(*) FROM cfdi")
    assert filas[0][0] == 5


def test_panel_se_pinta_sin_error(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app._pintar_panel()
    assert app.panel.winfo_children()          # tiene contenido


def test_ejecutar_sql_llena_tabla(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.txt_sql.delete("1.0", "end")
    app.txt_sql.insert("1.0", "SELECT tipo, COUNT(*) n FROM cfdi GROUP BY tipo")
    app._ejecutar_sql()
    assert len(app.tabla_sql.get_children()) == 2
    assert "filas" in app.v_sql_estado.get()


def test_sql_invalido_avisa(app, tmp_path, monkeypatch):
    import app as modulo
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    errores = []
    monkeypatch.setattr(modulo.messagebox, "showerror", lambda *a, **k: errores.append(a))
    app.txt_sql.delete("1.0", "end"); app.txt_sql.insert("1.0", "DELETE FROM cfdi")
    app._ejecutar_sql()
    assert errores


def test_cargar_consulta_ejemplo(app):
    app.v_consulta.set("Totales por mes")
    app._cargar_consulta_ejemplo()
    assert "GROUP BY" in app.txt_sql.get("1.0", "end")


def test_copiar_resultado_sql(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.txt_sql.delete("1.0", "end")
    app.txt_sql.insert("1.0", "SELECT tipo, total FROM cfdi")
    app._ejecutar_sql()
    app._copiar_sql()
    assert len(app.clipboard_get().split("\n")) >= 2


def test_exportar_sql_a_csv(app, tmp_path, monkeypatch):
    import app as modulo, csv as _csv
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.txt_sql.delete("1.0", "end"); app.txt_sql.insert("1.0", "SELECT tipo FROM cfdi")
    app._ejecutar_sql()
    destino = str(tmp_path / "q.csv")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_sql()
    with open(destino, encoding="utf-8-sig") as f:
        assert len(list(_csv.reader(f))) == 6


def test_columnas_configurables_afectan_copiado(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.cols_export = ["UUID", "Total"]
    app._copiar_tabla()
    assert app.clipboard_get().split("\n")[0] == "UUID\tTotal"


def test_export_csv_config_respeta_columnas(app, tmp_path, monkeypatch):
    import app as modulo, csv as _csv
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.cols_export = ["Fecha", "RFC Emisor", "Total"]
    destino = str(tmp_path / "cfg.csv")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_csv_config()
    with open(destino, encoding="utf-8-sig") as f:
        assert list(_csv.reader(f))[0] == ["Fecha", "RFC Emisor", "Total"]


def test_dialogo_columnas_preset_fiscal(app):
    import app as modulo
    capturado = {}
    dlg = modulo.DialogoColumnas(app, list(app.cols_export), lambda c: capturado.update(cols=c))
    dlg._preset_fiscal(); dlg._aplicar()
    assert "UUID" in capturado["cols"] and "Total" in capturado["cols"]
    assert "Archivo" not in capturado["cols"]       # el preset lo excluye


# ============================================================ visor de CFDI
def test_doble_clic_abre_visor(app, tmp_path):
    import app as modulo
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.tabla_datos.selection_set(app.tabla_datos.get_children()[0])
    app._abrir_visor()
    visores = [w for w in app.winfo_children() if isinstance(w, modulo.VisorCFDI)]
    assert visores
    for v in visores:
        v.destroy()


def test_visor_muestra_datos_del_cfdi(app, tmp_path):
    import app as modulo
    from sat_descarga.lector import leer_detalle
    carpeta = _poblar_xml(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta); app._leer_datos(esperar=True)
    archivo = app.filas_vista[0].datos["Archivo"]
    det = leer_detalle(archivo)
    v = modulo.VisorCFDI(app, det)
    v.update()
    assert v.det["uuid"]
    v.destroy()


def test_visor_genera_pdf(app, tmp_path, monkeypatch):
    import app as modulo
    from sat_descarga.lector import leer_detalle
    carpeta = _poblar_xml(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta); app._leer_datos(esperar=True)
    det = leer_detalle(app.filas_vista[0].datos["Archivo"])
    v = modulo.VisorCFDI(app, det)
    destino = str(tmp_path / "salida.pdf")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    monkeypatch.setattr(v, "_abrir_archivo", lambda r: None)
    v._pdf()
    from pathlib import Path as _P
    assert _P(destino).exists() and _P(destino).read_bytes()[:5] == b"%PDF-"
    v.destroy()


def test_visor_sin_seleccion_no_truena(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_xml(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.tabla_datos.selection_remove(app.tabla_datos.selection())
    app._abrir_visor()               # no debe reventar


# ============================================================ pagos, relacionados, builder
def _poblar_con_pago(carpeta_base):
    from test_lector import cfdi40, cfdi_pago
    rc = Path(carpeta_base) / "202608 HIHR" / "recibidos"; rc.mkdir(parents=True, exist_ok=True)
    em = Path(carpeta_base) / "202608 HIHR" / "emitidos"; em.mkdir(parents=True, exist_ok=True)
    for i in range(3):
        (em / f"E{i}.xml").write_bytes(cfdi40(uuid=f"1111111{i}-2222-3333-4444-555555555555"))
    (rc / "pago.xml").write_bytes(cfdi_pago())
    return str(carpeta_base)


def test_tabla_muestra_monto_pago(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_con_pago(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    pagos = [r for r in app.filas if r.datos["Comprobante"] == "Pago"]
    assert pagos and pagos[0].datos["Monto Total Pago"] == 2320.0
    assert pagos[0].datos["Imp Pagado"] == 2320.0


def test_visor_de_pago_tiene_datos(app, tmp_path):
    import app as modulo
    from sat_descarga.lector import leer_detalle
    carpeta = _poblar_con_pago(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta); app._leer_datos(esperar=True)
    pago = next(r for r in app.filas_vista if r.datos["Comprobante"] == "Pago")
    det = leer_detalle(pago.datos["Archivo"])
    v = modulo.VisorCFDI(app, det); v.update()
    assert v.det["pago_monto"] == 2320.0
    v.destroy()


def test_constructor_guiado_arma_y_ejecuta(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_con_pago(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.sub.select(2)
    app.gv_tipo.set("Emitido"); app.gv_campo.set("Total")
    app.gv_op.set("≥"); app.gv_valor.set("100")
    app._construir_consulta()
    sql = app.txt_sql.get("1.0", "end")
    assert "WHERE" in sql and "tipo = 'Emitido'" in sql
    assert len(app.tabla_sql.get_children()) >= 1


def test_constructor_sin_filtros_es_valido(app, tmp_path):
    app.v_carpeta_datos.set(_poblar_con_pago(tmp_path / "xmls"))
    app._leer_datos(esperar=True)
    app.gv_tipo.set("todos"); app.gv_valor.set("")
    app._construir_consulta()
    assert "SELECT" in app.txt_sql.get("1.0", "end")


def test_plantilla_pagos_se_carga(app):
    app.v_consulta.set("Pagos recibidos (REP) con detalle")
    app._cargar_consulta_ejemplo()
    assert "es_pago='Sí'" in app.txt_sql.get("1.0", "end")


def test_visor_tiene_volcado_completo(app, tmp_path):
    import app as modulo
    from sat_descarga.lector import leer_detalle, volcado_xml
    carpeta = _poblar_con_pago(tmp_path / "xmls")
    app.v_carpeta_datos.set(carpeta); app._leer_datos(esperar=True)
    pago = next(r for r in app.filas_vista if r.datos["Comprobante"] == "Pago")
    det = leer_detalle(pago.datos["Archivo"])
    # el volcado debe traer muchos campos
    assert len(volcado_xml(det["archivo"])) > 20
    v = modulo.VisorCFDI(app, det); v.update()
    v.destroy()


# ============================================================ v3: DESCARGA XML EDDP
import shutil as _sh

_REALES = Path(__file__).parent / "xml_reales"


def _cadena_real(app, tmp_path, con_duplicado=True):
    """Factura Kangoo + nota de crédito + REP (+ copia duplicada del REP)."""
    from sat_descarga.config import Perfil
    src = tmp_path / "cadena"; src.mkdir(exist_ok=True)
    for n in ("factura_kangoo_4237", "nc_anticipo_464", "pago_133k"):
        _sh.copy(_REALES / f"{n}.xml", src / f"{n}.xml")
    if con_duplicado:
        _sh.copy(_REALES / "pago_133k.xml", src / "pago_copia.xml")
    cer, key = _archivos_efirma(tmp_path, "FODA920130KD4")
    app.cfg.agregar(Perfil(rfc="FODA920130KD4", cer_path=cer, key_path=key,
                           key_password="p", nombre="ADRIANA"))
    app.v_carpeta_datos.set(str(src))
    app._leer_datos(esperar=True)
    return src


def test_nombre_de_la_app(app):
    assert app.title() == "DESCARGA XML EDDP"


def test_botones_de_exportar_estan_arriba_de_la_tabla(app):
    """Regresión del bug: en pantallas chicas la barra de abajo desaparecía."""
    app.deiconify(); app.geometry("1000x600"); app.nb.select(1); app.update()
    boton_excel = next(w for w in _todos(app.tab_datos)
                       if w.winfo_class() == "TButton" and "Excel" in str(w.cget("text")))
    assert boton_excel.winfo_rooty() < app.sub.winfo_rooty()
    assert boton_excel.winfo_ismapped()
    app.withdraw()


def _todos(w):
    yield w
    for h in w.winfo_children():
        yield from _todos(h)


def test_hay_boton_generar_pdfs(app):
    textos = [str(w.cget("text")) for w in _todos(app.tab_datos) if w.winfo_class() == "TButton"]
    assert any("Generar PDFs" in t for t in textos)


def test_cadena_real_en_tarjetas_e_inteligencia(app, tmp_path):
    _cadena_real(app, tmp_path)
    assert app._chips["cfdi"].cget("text") == "3"                       # 3 CFDI únicos
    assert "1 archivo duplicado ignorado" in app._chips_sub["cfdi"].cget("text")
    assert "45,000.00" in app._chips["saldo"].cget("text")               # saldo PPD real
    assert "1 parciales" in app.lbl_intel.cget("text")
    fac = next(r.datos for r in app.filas if r.datos["Folio"] == "4237")
    assert fac["Clase"] == "G" and fac["Estatus Pago"] == "Parcial"


def test_letras_por_clase_y_filtro_con_clic(app, tmp_path):
    _cadena_real(app, tmp_path)
    assert app._badges["G"].cget("text").split() == ["G", "1"]
    assert app._badges["P"].cget("text").split() == ["P", "1"]
    app._filtrar_clase("P")
    assert app.filas_vista and all(r.datos["Clase"] == "P" for r in app.filas_vista)
    app._filtrar_clase("P")                                               # segundo clic quita
    assert len(app.filas_vista) == len(app.filas)
    app._filtrar_clase("E"); app._limpiar_filtros()
    assert len(app.filas_vista) == len(app.filas)


def test_filas_coloreadas_por_clase(app, tmp_path):
    _cadena_real(app, tmp_path)
    tags = {app.tabla_datos.item(i, "tags")[-1] for i in app.tabla_datos.get_children()}
    assert {"cls_G", "cls_E", "cls_P"} <= tags


def test_generar_pdfs_desde_la_app(app, tmp_path, monkeypatch):
    import time as _t
    import app as modulo
    _cadena_real(app, tmp_path)
    destino = tmp_path / "pdfs"
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(modulo.filedialog, "askdirectory", lambda **k: str(destino))
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: avisos.append(a))
    monkeypatch.setattr(modulo.messagebox, "showwarning", lambda *a, **k: avisos.append(a))
    monkeypatch.setattr(app, "_abrir_ruta", lambda r: None)
    app.tabla_datos.selection_remove(app.tabla_datos.selection())
    app._generar_pdfs()
    limite = _t.time() + 30                       # esperar procesando eventos (como la app real)
    while not avisos and _t.time() < limite:
        app._vaciar_cola(); app.update(); _t.sleep(0.02)
    pdfs = sorted(p.name for p in destino.glob("*.pdf"))
    assert len(pdfs) == 3                                  # 4 filas, 3 CFDI, sin duplicado
    assert "PDF generados: 3" in avisos[0][1]
    assert app.cfg.opciones.ultima_carpeta_pdf == str(destino)


def test_generar_pdfs_solo_seleccion(app, tmp_path, monkeypatch):
    import time as _t
    import app as modulo
    _cadena_real(app, tmp_path)
    idx = next(i for i, r in enumerate(app.filas_vista) if r.datos["Folio"] == "4237")
    app.tabla_datos.selection_set(str(idx))
    destino = tmp_path / "sel"
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: True)
    monkeypatch.setattr(modulo.filedialog, "askdirectory", lambda **k: str(destino))
    hecho = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: hecho.append(a))
    monkeypatch.setattr(app, "_abrir_ruta", lambda r: None)
    app._generar_pdfs()
    limite = _t.time() + 30
    while not hecho and _t.time() < limite:
        app._vaciar_cola(); app.update(); _t.sleep(0.02)
    assert [p.name.split("_")[1] for p in destino.glob("*.pdf")] == ["G"]


def test_pdf_de_todos_ignora_la_seleccion(app, tmp_path, monkeypatch):
    import time as _t
    import app as modulo
    _cadena_real(app, tmp_path)
    idx = next(i for i, r in enumerate(app.filas_vista) if r.datos["Folio"] == "4237")
    app.tabla_datos.selection_set(str(idx))            # hay 1 seleccionado...
    destino = tmp_path / "todos"
    preguntas = []
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: preguntas.append(a) or True)
    monkeypatch.setattr(modulo.filedialog, "askdirectory", lambda **k: str(destino))
    hecho = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: hecho.append(a))
    monkeypatch.setattr(app, "_abrir_ruta", lambda r: None)
    app._generar_pdfs(todos=True)                      # ...pero "de todos" genera todo
    limite = _t.time() + 30
    while not hecho and _t.time() < limite:
        app._vaciar_cola(); app.update(); _t.sleep(0.02)
    assert len(list(destino.glob("*.pdf"))) == 3
    assert "Generar 3 PDF" in preguntas[0][1]
    assert not app.f_pdf.winfo_ismapped()


def test_visor_muestra_letra_y_seguimiento(app, tmp_path):
    import app as modulo
    _cadena_real(app, tmp_path)
    idx = next(i for i, r in enumerate(app.filas_vista) if r.datos["Folio"] == "4237")
    app.tabla_datos.selection_set(str(idx)); app._abrir_visor()
    v = [w for w in app.winfo_children() if isinstance(w, modulo.VisorCFDI)][-1]
    v.update()
    textos = [str(w.cget("text")) for w in _todos(v) if w.winfo_class() == "Label"]
    assert "G" in textos and "GASTO" in textos
    assert "Parcial" in textos
    assert any("45,000.00" in t for t in textos)
    v.destroy()


def test_visor_genera_pdf_con_config(app, tmp_path, monkeypatch):
    import app as modulo
    _cadena_real(app, tmp_path)
    idx = next(i for i, r in enumerate(app.filas_vista) if r.datos["Folio"] == "4237")
    fila = app.filas_vista[idx].datos
    v = modulo.VisorCFDI(app, modulo.leer_detalle(fila["Archivo"]), fila)
    destino = str(tmp_path / "uno.pdf")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    monkeypatch.setattr(v, "_abrir_archivo", lambda r: None)
    v._pdf()
    from conftest import texto_pdf
    txt = texto_pdf(destino)
    assert "GASTO" in txt and "Parcial" in txt
    v.destroy()


def test_dialogo_config_pdf_guarda(app, tmp_path):
    """Antes: diálogo 'general + por RFC' en pdf_config.json. Ahora (Fase 2) el
    mismo botón edita el diseño general o el perfil del cliente."""
    import app as modulo
    from sat_descarga.perfiles import Perfiles
    _cadena_real(app, tmp_path, con_duplicado=False)
    dlg = modulo.DialogoPerfilReporte(app, None)               # sin perfil = diseño general
    dlg.v["pie"].set("EDDP Servicios · erickddp.com")
    dlg.v["color_primario"].set("#123456")
    dlg.b["sellos"].set(True)
    dlg._guardar()
    nuevo = app.perfiles_rep.crear("ADRIANA", rfcs=["FODA920130KD4"])
    dlg = modulo.DialogoPerfilReporte(app, nuevo)
    dlg.v["titulo"].set("ADRIANA G. FONSECA")
    dlg._guardar()
    otra = Perfiles.cargar(str(app.perfiles_rep.ruta))
    assert otra.general["pie"] == "EDDP Servicios · erickddp.com"
    assert otra.general["secciones"]["sellos"] is True
    assert otra.por_rfc("FODA920130KD4").diseno["titulo"] == "ADRIANA G. FONSECA"


def test_respaldar_y_restaurar_desde_la_app(app, tmp_path, monkeypatch):
    import app as modulo
    _cadena_real(app, tmp_path, con_duplicado=False)
    zip_ = str(tmp_path / "respaldo.zip")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: zip_)
    avisos = []
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: avisos.append(a[1]))
    app._respaldar()
    assert Path(zip_).exists() and "Contiene tus llaves" in avisos[-1]

    app.cfg.perfiles.clear(); app.cfg.guardar(); app._refrescar_lista()
    monkeypatch.setattr(modulo.filedialog, "askopenfilename", lambda **k: zip_)
    monkeypatch.setattr(modulo.messagebox, "askyesno", lambda *a, **k: True)
    app._restaurar()
    assert app.cfg.por_rfc("FODA920130KD4") is not None
    assert len(app.tabla.get_children()) == 1 and "1 RFC nuevos" in avisos[-1]


def test_zoom_se_aplica_y_se_recuerda(app):
    from sat_descarga.config import Config
    base = app.tema.fuente("base").cget("size")
    app._zoom(0.2)
    assert app.tema.fuente("base").cget("size") > base
    assert Config.desde_archivo(app.cfg.ruta).opciones.zoom == 1.2
    app._zoom(None)
    assert app.tema.fuente("base").cget("size") == base


def test_memoria_entre_sesiones_desde_la_app(app, tmp_path):
    """Lee el REP solo; luego la factura sola en otra carpeta: encuentra su pago."""
    from sat_descarga.config import Perfil
    cer, key = _archivos_efirma(tmp_path, "FODA920130KD4")
    app.cfg.agregar(Perfil(rfc="FODA920130KD4", cer_path=cer, key_path=key, key_password="p", nombre="A"))
    a = tmp_path / "solo_pago"; a.mkdir(); _sh.copy(_REALES / "pago_133k.xml", a / "p.xml")
    b = tmp_path / "solo_factura"; b.mkdir(); _sh.copy(_REALES / "factura_kangoo_4237.xml", b / "f.xml")
    app.v_carpeta_datos.set(str(a)); app._leer_datos(esperar=True)
    app.v_carpeta_datos.set(str(b)); app._leer_datos(esperar=True)
    fac = app.filas[0].datos
    assert fac["Pagado REP"] == 124389.04 and fac["Estatus Pago"] == "Parcial"


# ============================================================ formato exacto de copiado
FORMATO_RICK = ("Tipo\tFecha\tAño\tMes\tDía\tHora\tComprobante\tEfecto\tVersión\tSerie\tFolio\t"
    "RFC Emisor\tNombre Emisor\tRégimen Emisor\tRFC Receptor\tNombre Receptor\tRégimen Receptor\t"
    "Uso CFDI\tDomicilio Receptor\tSubTotal\tDescuento\tIVA Trasladado\tIEPS Trasladado\t"
    "IVA Retenido\tISR Retenido\tRet. Total\tTotal\tBase IVA 16\tBase IVA 8\tBase IVA 0\t"
    "Base Exento\tMoneda\tTipo Cambio\tMétodo Pago\tForma Pago\tCondiciones Pago\tExportación\t"
    "Es Pago\tMonto Total Pago\tFecha Pago\tForma Pago P\tMoneda Pago\tTC Pago\tNum Operación\t"
    "Doc Pagado UUID\tDoc Serie\tDoc Folio\tParcialidad\tSaldo Anterior\tImp Pagado\t"
    "Saldo Insoluto\tBase Doc\tIVA Doc\tTasa Doc\tObjeto Imp Doc\tRelacionados\tTipo Relación\t"
    "Conceptos\tDescripción\tClaves ProdServ\tNo Identificación\tComplementos\tUUID\t"
    "Fecha Timbrado\tRFC PAC")


def test_copiar_tabla_sale_exactamente_con_el_formato_de_rick(app, tmp_path):
    _cadena_real(app, tmp_path, con_duplicado=False)
    app._copiar_tabla()
    lineas = app.clipboard_get().split("\n")
    assert lineas[0] == FORMATO_RICK                     # ni un campo más ni menos
    assert all(len(l.split("\t")) == 65 for l in lineas)
    assert "Clase" not in lineas[0] and "Archivo" not in lineas[0]


def test_excel_conserva_todas_las_columnas(app, tmp_path, monkeypatch):
    import app as modulo
    from openpyxl import load_workbook
    from sat_descarga.lector import COLUMNAS
    _cadena_real(app, tmp_path, con_duplicado=False)
    destino = str(tmp_path / "todo.xlsx")
    monkeypatch.setattr(modulo.filedialog, "asksaveasfilename", lambda **k: destino)
    monkeypatch.setattr(modulo.messagebox, "showinfo", lambda *a, **k: None)
    app._exportar_excel()
    assert [c.value for c in load_workbook(destino)["CFDI"][1]] == COLUMNAS   # el otro reporte


def test_dialogo_columnas_cabe_en_laptop(app):
    import app as modulo
    dlg = modulo.DialogoColumnas(app, list(app.cols_export), lambda c: None)
    dlg.update_idletasks()
    assert dlg.winfo_reqheight() < 700 and dlg.winfo_reqwidth() < 1366
    textos = [str(w.cget("text")) for w in _todos(dlg) if w.winfo_class() == "TButton"]
    assert "Guardar" in textos
    assert dlg.v_cuenta.get() == "65 de 73 seleccionadas"
    dlg._aplicar()


def test_cerrar_con_la_x_guarda_y_sobrevive_a_reabrir(app):
    import app as modulo
    from sat_descarga.config import Config
    dlg = modulo.DialogoColumnas(app, list(app.cols_export), app._set_columnas)
    dlg._marcar(False)
    for c in ("Fecha", "RFC Emisor", "Total", "UUID"):
        dlg.vars[c].set(True)
    dlg.tk.call(dlg.protocol("WM_DELETE_WINDOW"))          # lo mismo que la X
    assert app.cols_export == ["Fecha", "RFC Emisor", "Total", "UUID"]
    assert Config.desde_archivo(app.cfg.ruta).opciones.columnas_copiado == app.cols_export
    otra = modulo.App(ruta_config=app.cfg.ruta); otra.withdraw()
    assert otra.cols_export == ["Fecha", "RFC Emisor", "Total", "UUID"]
    otra.destroy()


def test_volver_al_formato_eddp(app):
    import app as modulo
    from sat_descarga.lector import COLUMNAS_COPIADO
    app._set_columnas(["UUID"])
    dlg = modulo.DialogoColumnas(app, list(app.cols_export), app._set_columnas)
    dlg._preset_eddp(); dlg._aplicar()
    assert app.cols_export == list(COLUMNAS_COPIADO)
    assert app.cfg.opciones.columnas_copiado == []          # vacío = formato EDDP


# ============================================================ marca y fluidez
def test_la_app_trae_logo_e_icono(app):
    recursos = Path(__file__).parent.parent / "recursos"
    for nombre in ("icono.ico", "icono.png", "logo_blanco.png"):
        assert (recursos / nombre).exists(), nombre
    assert app._logo is not None and app._icono is not None


def test_buscador_espera_a_que_dejes_de_teclear(app, tmp_path):
    import time as _t
    _cadena_real(app, tmp_path)
    pintadas = []
    real = app._pintar_tabla
    app._pintar_tabla = lambda: pintadas.append(1) or real()
    for letra in "4237":
        app.v_buscar.set(app.v_buscar.get() + letra)
        app._filtrar_pronto()
    assert pintadas == []                               # aún no repinta
    limite = _t.time() + 3
    while not pintadas and _t.time() < limite:
        app.update(); _t.sleep(0.02)
    assert len(pintadas) == 1                           # una sola vez, al final
    assert {r.datos["Folio"] for r in app.filas_vista} == {"4237"}
    del app._pintar_tabla


def test_buscador_filtra_por_fecha(app, tmp_path):
    _cadena_real(app, tmp_path)
    app.v_buscar.set("2026-08-12"); app._aplicar_filtros()
    assert app.filas_vista and all(str(r.datos["Fecha"]).startswith("2026-08-12") for r in app.filas_vista)
